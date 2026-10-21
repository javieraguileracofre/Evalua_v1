# integrations/sii/parser.py
# -*- coding: utf-8 -*-
"""Lee el archivo descargado. No escribe en la base de datos."""
from __future__ import annotations

import csv
import io
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from integrations.sii.errors import EstructuraInvalida

_ALIASES: dict[str, tuple[str, ...]] = {
    "tipo_doc": ("tipo doc", "tipo documento", "tipo dte", "tipo docto"),
    "folio": ("folio",),
    "rut_proveedor": ("rut proveedor",),
    "rut_cliente": ("rut cliente", "rut receptor"),
    "razon": ("razon social",),
    "fecha_emision": ("fecha docto", "fecha documento", "fecha emision"),
    "fecha_recepcion": ("fecha recepcion",),
    "exento": ("monto exento",),
    "neto": ("monto neto",),
    "neto_af": ("monto neto activo fijo",),
    "iva": ("monto iva",),
    "iva_rec": ("monto iva recuperable",),
    "iva_no_rec": ("monto iva no recuperable",),
    "iva_af": ("iva activo fijo",),
    "iva_uso": ("iva uso comun",),
    "total": ("monto total",),
    "otro": ("valor otro impuesto", "valor otro imp.", "valor otro imp"),
    "estado": ("estado",),
}


@dataclass(frozen=True)
class TablaSii:
    columnas: dict[str, str]
    filas: list[dict[str, str]]


def _sin_acentos(value: str) -> str:
    base = unicodedata.normalize("NFD", value or "")
    sin = "".join(ch for ch in base if unicodedata.category(ch) != "Mn")
    return " ".join(sin.lower().replace(".", " ").split())


def mapear_encabezados(encabezados: list[str]) -> dict[str, str]:
    normalizados = {_sin_acentos(col): col for col in encabezados if _sin_acentos(col)}
    encontrados: dict[str, str] = {}
    for canon, aliases in _ALIASES.items():
        for alias in aliases:
            clave = _sin_acentos(alias)
            if clave in normalizados:
                encontrados[canon] = normalizados[clave]
                break
    if "tipo_doc" not in encontrados or "folio" not in encontrados:
        raise EstructuraInvalida(
            "El archivo no tiene las columnas Tipo Doc y Folio del Registro de Compras y Ventas."
        )
    if "neto" not in encontrados and "total" not in encontrados:
        raise EstructuraInvalida(
            "El archivo no trae montos reconocibles (Monto Neto o Monto Total)."
        )
    return encontrados


def _decodificar(raw: bytes) -> str:
    muestra = raw.lstrip()[:300].lower()
    if muestra.startswith(b"<") or muestra.startswith(b"<!doctype") or b"<html" in muestra:
        raise EstructuraInvalida(
            "El archivo no es un libro del SII: el sitio devolvió una página web."
        )
    for codec in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(codec)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def _leer_csv(texto: str) -> tuple[list[str], list[list[str]]]:
    muestra = texto[:4000]
    delimitador = ";" if muestra.count(";") >= muestra.count(",") else ","
    lector = csv.reader(io.StringIO(texto), delimiter=delimitador)
    filas = [fila for fila in lector if any(celda.strip() for celda in fila)]
    if not filas:
        raise EstructuraInvalida("El archivo del SII está vacío.")
    return [celda.strip() for celda in filas[0]], filas[1:]


def _leer_xlsx(raw: bytes) -> tuple[list[str], list[list[str]]]:
    from openpyxl import load_workbook

    libro = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        hoja = libro.active
        if hoja is None:
            raise EstructuraInvalida("El archivo Excel del SII no tiene hojas.")
        filas: list[list[str]] = []
        for row in hoja.iter_rows(values_only=True):
            celdas = ["" if value is None else str(value).strip() for value in row]
            if any(celdas):
                filas.append(celdas)
    finally:
        libro.close()
    if not filas:
        raise EstructuraInvalida("El archivo Excel del SII está vacío.")
    return filas[0], filas[1:]


def leer_tabla(path: Path) -> TablaSii:
    raw = path.read_bytes()
    if not raw.strip():
        raise EstructuraInvalida("El archivo del SII está vacío.")
    if raw[:2] == b"PK" or path.suffix.lower() in {".xlsx", ".xlsm"}:
        encabezados, cuerpo = _leer_xlsx(raw)
    else:
        encabezados, cuerpo = _leer_csv(_decodificar(raw))
    columnas = mapear_encabezados(encabezados)
    filas: list[dict[str, str]] = []
    for numero, celdas in enumerate(cuerpo, start=2):
        ancho = max(len(encabezados), len(celdas))
        pares = {
            encabezados[i]: (celdas[i].strip() if i < len(celdas) else "")
            for i in range(ancho)
            if i < len(encabezados) and encabezados[i]
        }
        fila = {canon: pares.get(real, "") for canon, real in columnas.items()}
        fila["_fila"] = str(numero)
        filas.append(fila)
    return TablaSii(columnas=columnas, filas=filas)
