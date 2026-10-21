# -*- coding: utf-8 -*-
"""Fase 1 cotización leasing financiero: cálculo, estados, export y permisos."""
from __future__ import annotations

from datetime import date
from decimal import Decimal
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from starlette.requests import Request

from core.rbac import (
    guard_leasing_fin_consulta,
    guard_leasing_fin_mutacion,
    require_leasing_fin_consulta_api,
    require_leasing_fin_mutacion_api,
)
from schemas.comercial.leasing_cotizacion import LeasingSimulacionInput
from services.leasing_financiero import (
    calcular_tabla_amortizacion,
    tasa_efectiva_periodo,
    simular_cotizacion,
)
from services.leasing_financiero_estados import estado_para_db, puede_comprar_activo
from services.leasing_financiero_export import build_amortizacion_excel, nombre_archivo_export
from services.leasing_financiero_cotizacion_pdf import build_cotizador_pdf
from services.leasing_financiero_workflow import exigir_acuerdo_para_compra, puede_transicionar


def _cot(**kwargs):
    base = dict(
        plazo=12,
        monto_financiado=Decimal("1000000"),
        valor_neto=Decimal("1000000"),
        monto=None,
        tasa=Decimal("0.12"),
        tasa_tipo="ANUAL",
        periodos_gracia=0,
        periodicidad="MENSUAL",
        opcion_compra=Decimal("0"),
        fecha_inicio=date(2026, 1, 15),
        fecha_primera_cuota=None,
        iva_aplica=False,
        iva_tasa=None,
    )
    base.update(kwargs)
    return SimpleNamespace(**base)


def test_tasa_anual_a_mensual_efectiva():
    i = tasa_efectiva_periodo(Decimal("0.12"), tasa_tipo="ANUAL", periodicidad="MENSUAL")
    # (1.12)^(1/12) - 1 ≈ 0.0094888
    assert i >= Decimal("0.0094")
    assert i <= Decimal("0.0095")


def test_tasa_mensual_se_usa_directa():
    i = tasa_efectiva_periodo(Decimal("0.01"), tasa_tipo="MENSUAL", periodicidad="MENSUAL")
    assert i == Decimal("0.0100")


def test_frances_cuota_conocida_tasa_mensual():
    """PV=1.000.000, n=12, i=1% mensual, FV=0 → PMT ≈ 88.848,79."""
    cot = _cot(tasa=Decimal("0.01"), tasa_tipo="MENSUAL", opcion_compra=Decimal("0"))
    tabla = calcular_tabla_amortizacion(cot)
    rentas = [c for c in tabla if not c.es_gracia and not c.es_opcion_compra]
    assert len(rentas) == 12
    assert rentas[0].cuota == Decimal("88848.79")
    assert tabla[-1].saldo_final == Decimal("0.00")
    suma_amort = sum((c.amortizacion for c in tabla), Decimal("0"))
    assert suma_amort == Decimal("1000000.00")


def test_tabla_incluye_iva_y_total_cuota():
    cot = _cot(iva_aplica=True, iva_tasa=Decimal("19"), tasa_tipo="MENSUAL", tasa=Decimal("0.01"))
    tabla = calcular_tabla_amortizacion(cot)
    cuota = tabla[0]
    assert cuota.iva == Decimal("16881.27")
    assert cuota.total_cuota == Decimal("105730.06")


def test_rechaza_valores_negativos_y_pie_mayor():
    with pytest.raises(ValueError, match="negativ"):
        calcular_tabla_amortizacion(_cot(opcion_compra=Decimal("-1")))
    with pytest.raises(ValueError, match="plazo"):
        calcular_tabla_amortizacion(_cot(plazo=0))
    from services.leasing_financiero import calcular_pago_inicial

    with pytest.raises(ValueError, match="superar"):
        calcular_pago_inicial(Decimal("1000"), "MONTO", Decimal("1500"))


def test_flujo_estados_fase1():
    assert puede_transicionar("BORRADOR", "COTIZADA")
    assert puede_transicionar("COTIZADA", "EN_EVALUACION")
    assert puede_transicionar("EN_EVALUACION", "APROBADA")
    assert puede_transicionar("APROBADA", "ACUERDO_FIRMADO")
    assert puede_transicionar("ACUERDO_FIRMADO", "ACTIVO_PENDIENTE_COMPRA")
    assert puede_transicionar("ACTIVO_PENDIENTE_COMPRA", "ACTIVO_ADQUIRIDO")
    assert not puede_transicionar("COTIZADA", "ACTIVO_ADQUIRIDO")


def test_estado_para_db_cumple_check_legacy():
    assert estado_para_db("EN_EVALUACION") == "EN_ANALISIS_CREDITO"
    assert estado_para_db("ACUERDO_FIRMADO") == "EN_FORMALIZACION"
    assert estado_para_db("ACTIVO_PENDIENTE_COMPRA") == "DOCUMENTACION_COMPLETA"
    assert estado_para_db("ACTIVO_ADQUIRIDO") == "ACTIVADA"
    assert estado_para_db("APROBADA") == "APROBADA"
    assert estado_para_db("COTIZADA") == "COTIZADA"


def test_modelo_persiste_estado_compatible_con_check():
    from models.comercial.leasing_financiero_cotizacion import LeasingFinancieroCotizacion

    cot = LeasingFinancieroCotizacion(cliente_id=1, estado="EN_EVALUACION")
    assert cot.estado == "EN_EVALUACION"
    assert cot.estado_db == "EN_ANALISIS_CREDITO"
    assert cot.workflow_json.get("estado_canonico") == "EN_EVALUACION"


def test_compra_activo_solo_despues_de_acuerdo():
    assert not puede_comprar_activo("BORRADOR")
    assert not puede_comprar_activo("COTIZADA")
    assert not puede_comprar_activo("EN_EVALUACION")
    assert not puede_comprar_activo("APROBADA")
    assert puede_comprar_activo("ACUERDO_FIRMADO")
    assert puede_comprar_activo("ACTIVO_PENDIENTE_COMPRA")
    with pytest.raises(ValueError, match="acuerdo"):
        exigir_acuerdo_para_compra("COTIZADA")
    exigir_acuerdo_para_compra("ACUERDO_FIRMADO")


def test_excel_corporativo_dos_hojas_y_nombre():
    cot = _cot(id=7, moneda="CLP", estado="COTIZADA", bien_descripcion="Camión")
    cot.cliente = SimpleNamespace(razon_social="Cliente Demo SpA", rut="76.111.111-1")
    cot.proveedor = None
    cot.concesionario = "Volvo"
    cot.tasa_tipo = "ANUAL"
    cot.gastos_administrativos = Decimal("0")
    cot.comision_apertura = Decimal("0")
    cot.seguro_monto_uf = Decimal("0")
    cot.observaciones = "Demo"
    cot.ejecutivo = "Ana"
    tabla = calcular_tabla_amortizacion(cot)
    bio = build_amortizacion_excel(cot, tabla)
    from openpyxl import load_workbook

    wb = load_workbook(bio)
    assert "Resumen" in wb.sheetnames
    assert "Amortización" in wb.sheetnames
    assert wb["Amortización"]["A3"].value == "N°"
    assert nombre_archivo_export(cot, "xlsx") == "leasing_financiero_cotizacion_7_Cliente_Demo_SpA.xlsx"


def test_pdf_corporativo_completo():
    cot = _cot(id=8, moneda="CLP", estado="COTIZADA", bien_descripcion="Grúa")
    cot.cliente = SimpleNamespace(razon_social="Beta Ltda", rut="76.222.222-2", email="a@b.cl")
    cot.proveedor = None
    cot.concesionario = None
    cot.activo = None
    cot.observaciones = "Uso interno"
    cot.ejecutivo = "Luis"
    cot.gastos_administrativos = Decimal("10000")
    cot.comision_apertura = Decimal("0")
    cot.seguro_monto_uf = Decimal("0")
    res = simular_cotizacion(
        LeasingSimulacionInput(
            moneda="CLP",
            valor_neto=Decimal("1000000"),
            tasa=Decimal("0.12"),
            tasa_tipo="ANUAL",
            plazo=12,
        )
    )
    tabla = calcular_tabla_amortizacion(cot)
    pdf = build_cotizador_pdf(cotizacion=cot, resumen=res, tabla=tabla)
    assert pdf.startswith(b"%PDF")
    assert len(pdf) > 2000


def _req(roles, modules, method="GET"):
    scope = {"type": "http", "method": method, "path": "/", "headers": []}
    req = Request(scope)
    req.state.auth_user = {"roles": roles, "visibleModules": modules}
    return req


def test_permiso_consulta_bloquea_rrhh():
    req = _req(["RRHH"], ["RRHH", "PRINCIPAL"])
    redir = guard_leasing_fin_consulta(req)
    assert redir is not None
    assert redir.status_code == 303


def test_permiso_mutacion_bloquea_consulta():
    req = _req(["CONSULTA"], ["COMERCIAL", "PRINCIPAL"], method="POST")
    redir = guard_leasing_fin_mutacion(req)
    assert redir is not None


def test_api_consulta_sin_permiso_403():
    req = _req(["RRHH"], ["RRHH"])
    with pytest.raises(HTTPException) as err:
        require_leasing_fin_consulta_api(req)
    assert err.value.status_code == 403


def test_api_mutacion_sin_permiso_403():
    req = _req(["CONSULTA"], ["COMERCIAL"], method="POST")
    with pytest.raises(HTTPException) as err:
        require_leasing_fin_mutacion_api(req)
    assert err.value.status_code == 403
